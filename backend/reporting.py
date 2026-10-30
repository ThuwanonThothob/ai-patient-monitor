"""Reporting helpers for the Smart Ward dashboard.

The generated report is a decision-support handoff, not a clinical order.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any


def format_patient_handoff_text(
    patient_id: str,
    profile: dict[str, Any],
    history_rows: list[dict[str, Any]],
    assessment: dict[str, Any],
) -> str:
    """Build a portable, patient-scoped text handoff from synthetic data."""
    latest = history_rows[0] if history_rows else {}
    stats = assessment.get("stats", {})
    diagnoses = ", ".join(profile.get("diagnoses", [profile.get("condition", "-")]))
    lines = [
        "AI Smart Ward — Patient Handoff (synthetic demonstration data)",
        f"Generated: {datetime.now().astimezone().strftime('%Y-%m-%d %H:%M:%S %Z')}",
        "",
        "Patient location",
        f"Patient ID: {patient_id}",
        f"Zone / Room / Bed: {profile.get('zone', '-')} / {profile.get('room_number', '-')} / {profile.get('bed_number', '-')}",
        f"Name in demo: {profile.get('name', '-')}",
        f"Age / Sex / Blood group: {profile.get('age', '-')} / {profile.get('gender', '-')} / {profile.get('blood_group', '-')}",
        f"Assigned clinician (demo): {profile.get('doctor', '-')}",
        "",
        "Care profile",
        f"Condition: {profile.get('condition', '-')}",
        f"Diagnoses / monitoring focus: {diagnoses}",
        f"Mobility: {profile.get('mobility', '-')}",
        f"Baseline respiratory rate: {profile.get('respiratory_rate', '-')} breaths/min",
        f"Sleep: {profile.get('sleep_hours', '-')} hours ({profile.get('sleep_quality', '-')})",
        f"Blood glucose: {profile.get('blood_glucose_mg_dl', '-')} mg/dL",
        f"Cholesterol: {profile.get('cholesterol_mg_dl', '-')} mg/dL",
        f"ECG rhythm: {profile.get('ecg_rhythm', '-')}",
        "",
        "Latest monitored reading",
        f"Heart rate: {latest.get('heart_rate', '-')} bpm | SpO2: {latest.get('spo2', '-')}% | RR: {latest.get('respiratory_rate', '-')} breaths/min",
        f"Blood glucose: {latest.get('blood_glucose', '-')} mg/dL | ECG: {latest.get('ecg_rhythm', '-')}",
        f"Monitor status: {latest.get('status', 'No reading')} | Simulator state: {latest.get('simulation_state', '-')}",
        "",
        "AI-assisted review",
        f"Demo rule priority (not probability): {assessment.get('risk_score', '—')} points — {assessment.get('risk_level', '-')}",
        f"HR average / maximum: {stats.get('avg_hr', '-')} / {stats.get('max_hr', '-')} bpm",
        f"SpO2 average / minimum: {stats.get('avg_spo2', '-')}% / {stats.get('min_spo2', '-')}%",
        assessment.get("summary", "No assessment available"),
        "Recommended review: " + ("; ".join(assessment.get("recommendations", [])) or "Continue routine observation"),
        "",
        "Safety note: This is synthetic demonstration data and decision support only. A qualified person must verify readings and approve all care, contact, or transfer decisions.",
    ]
    return "\n".join(lines)


def create_patient_handoff_docx(
    patient_id: str,
    profile: dict[str, Any],
    history_rows: list[dict[str, Any]],
    assessment: dict[str, Any],
) -> bytes:
    """Create an English DOCX handoff because the local renderer lacks Thai glyphs."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as error:
        raise RuntimeError("python-docx is not installed; run pip install -r requirements.txt") from error

    latest = history_rows[0] if history_rows else {}
    stats = assessment.get("stats", {})
    diagnoses = ", ".join(profile.get("diagnoses", [profile.get("condition", "-")]))
    document = Document()
    settings = document.sections[0]
    settings.top_margin = Inches(0.65)
    settings.bottom_margin = Inches(0.65)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Patient Handoff Report")
    title_run.bold = True
    title_run.font.name = "Tahoma"
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = document.add_paragraph("AI Smart Ward — synthetic demonstration record")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_table(title_text: str, values: list[tuple[str, Any]]) -> None:
        document.add_heading(title_text, level=1)
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in values:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = str(value if value not in (None, "") else "—")

    add_table("Patient location", [
        ("Patient ID", patient_id),
        ("Zone", profile.get("zone")),
        ("Room / bed", f"{profile.get('room_number', '—')} / {profile.get('bed_number', '—')}"),
        ("Name in demo", profile.get("name")),
        ("Age / sex / blood group", f"{profile.get('age', '—')} / {profile.get('gender', '—')} / {profile.get('blood_group', '—')}"),
        ("Assigned clinician (demo)", profile.get("doctor")),
    ])
    add_table("Care profile", [
        ("Condition", profile.get("condition")),
        ("Diagnoses / monitoring focus", diagnoses),
        ("Mobility", profile.get("mobility")),
        ("Baseline respiratory rate", f"{profile.get('respiratory_rate', '—')} breaths/min"),
        ("Sleep", f"{profile.get('sleep_hours', '—')} hours ({profile.get('sleep_quality', '—')})"),
        ("Blood glucose / cholesterol", f"{profile.get('blood_glucose_mg_dl', '—')} / {profile.get('cholesterol_mg_dl', '—')} mg/dL"),
        ("ECG rhythm", profile.get("ecg_rhythm")),
    ])
    add_table("Latest monitored reading", [
        ("Heart rate / SpO2", f"{latest.get('heart_rate', '—')} bpm / {latest.get('spo2', '—')}%"),
        ("Respiratory rate", f"{latest.get('respiratory_rate', '—')} breaths/min"),
        ("Blood glucose / ECG", f"{latest.get('blood_glucose', '—')} mg/dL / {latest.get('ecg_rhythm', '—')}"),
        ("Monitor state", latest.get("status", "No reading")),
        ("Simulation state", latest.get("simulation_state", "—")),
    ])
    add_table("AI-assisted review", [
        ("Demo rule priority (not probability)", f"{assessment.get('risk_score', '—')} points — {assessment.get('risk_level', '—')}"),
        ("HR average / maximum", f"{stats.get('avg_hr', '—')} / {stats.get('max_hr', '—')} bpm"),
        ("SpO2 average / minimum", f"{stats.get('avg_spo2', '—')}% / {stats.get('min_spo2', '—')}%"),
        ("Review instruction", "; ".join(assessment.get("recommendations", [])) or "Continue routine observation"),
    ])
    safety = document.add_paragraph(
        "Safety note: This report contains synthetic demonstration data and is decision support only. "
        "A qualified person must verify readings and approve every care, contact, or transfer decision."
    )
    safety.italic = True
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Tahoma"
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
    for report_table in document.tables:
        for row in report_table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Tahoma"
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(0, 0, 0)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def format_dashboard_text(section: str, snapshot: dict[str, Any], ai_summary: str) -> str:
    lines = [
        "AI Smart Ward System Report",
        f"Section: {section.title()}",
        f"Generated: {snapshot['generated_at']}",
        "",
        "Operational summary",
        f"Patients with readings: {snapshot['patient_count']}",
        f"Vital readings in database: {snapshot['reading_count']}",
        f"Critical readings: {snapshot['critical_count']}",
        f"Warning readings: {snapshot['warning_count']}",
        "",
        "AI decision support summary",
        ai_summary,
        "",
        "Patient snapshot",
    ]
    for patient in snapshot["patients"]:
        lines.append(
            f"{patient['patient_id']} | {patient['name']} | "
            f"HR avg {patient['avg_hr']} bpm | SpO2 min {patient['min_spo2']}% | "
            f"critical {patient['critical_count']} | warning {patient['warning_count']}"
        )
    lines.extend([
        "",
        "This report is decision support only. Verify readings and obtain clinical review before action.",
    ])
    return "\n".join(lines)


def create_report_docx(section: str, snapshot: dict[str, Any], ai_summary: str) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError as error:
        raise RuntimeError("python-docx is not installed; run pip install -r requirements.txt") from error

    document = Document()
    section_settings = document.sections[0]
    section_settings.top_margin = Inches(0.7)
    section_settings.bottom_margin = Inches(0.7)

    title = document.add_heading("AI Smart Ward System Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(f"{section.title()} report  Generated {snapshot['generated_at']}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("Operational summary", level=1)
    for label, value in (
        ("Patients with readings", snapshot["patient_count"]),
        ("Vital readings in database", snapshot["reading_count"]),
        ("Critical readings", snapshot["critical_count"]),
        ("Warning readings", snapshot["warning_count"]),
    ):
        document.add_paragraph(f"{label}: {value}")

    document.add_heading("AI decision support summary", level=1)
    priority_beds = [patient["patient_id"] for patient in snapshot["patients"] if patient["critical_count"] or patient["warning_count"]]
    document.add_paragraph(
        "AI-assisted review found "
        f"{snapshot['critical_count']} critical and {snapshot['warning_count']} warning readings. "
        f"Review priority: {', '.join(priority_beds) if priority_beds else 'no flagged beds'}. "
        "The Thai AI narrative is available in the accompanying Text report."
    )

    document.add_heading("Patient snapshot", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headings = ["Bed", "Patient", "HR average", "SpO2 minimum", "Critical", "Warning"]
    for cell, heading in zip(table.rows[0].cells, headings):
        cell.text = heading
    for patient in snapshot["patients"]:
        values = [
            patient["patient_id"], f"Bed {patient['patient_id']} record", str(patient["avg_hr"]),
            str(patient["min_spo2"]), str(patient["critical_count"]), str(patient["warning_count"]),
        ]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value

    document.add_paragraph(
        "This report is decision support only. Verify readings and obtain clinical review before action."
    ).italic = True
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Tahoma"
            run.font.size = Pt(10.5)
    for report_table in document.tables:
        for row in report_table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Tahoma"
                        run.font.size = Pt(9.5)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
